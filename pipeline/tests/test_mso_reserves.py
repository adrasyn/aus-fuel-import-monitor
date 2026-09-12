import json
from datetime import datetime

import pytest
from openpyxl import Workbook

from pipeline.mso_reserves import parse_latest_reserves, build_reserves_json


HEADER = (
    "Obligation Date",
    "Stocks required under MSO (ML)",
    "Stock held under MSO (ML)",
    "Stock held under MSO (Days equivalent) [2]",
)


def _make_workbook(path, sheets):
    """sheets: {sheet_name: [(date, *values..., days), ...]} — mirrors the real
    DCCEEW layout: header row, fuel-name row, then one row per obligation date.
    Gasoline/Diesel carry an extra s16A column; Kerosene doesn't."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Title and index").append(["Minimum Stockholding Obligation - Weekly Snapshot"])
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(HEADER)
        ws.append((None, name, name, name))
        for row in rows:
            ws.append(row)
    wb.save(path)
    return str(path)


def _standard_sheets(latest=datetime(2026, 9, 1)):
    return {
        "Gasoline": [
            (datetime(2026, 8, 25), 1067, 862, 1743, 41),
            (latest, 1067, 862, 1823, 43),
        ],
        "Diesel": [
            (datetime(2026, 8, 25), 2742, 2225, 3351, 36),
            (latest, 2742, 2225, 3039, 33),
        ],
        "Kerosene": [
            (datetime(2026, 8, 25), 663, 840, 31),
            (latest, 663, 901, 33),
        ],
    }


def test_parse_latest_reserves_reads_last_row_of_each_sheet(tmp_path):
    path = _make_workbook(tmp_path / "mso.xlsx", _standard_sheets())

    result = parse_latest_reserves(path)

    assert result["as_of"] == "2026-09-01"
    assert result["fuels"] == [
        {"key": "petrol", "label": "Petrol", "days": 43},
        {"key": "kerosene", "label": "Kerosene", "days": 33},
        {"key": "diesel", "label": "Diesel", "days": 33},
    ]


def test_parse_latest_reserves_ignores_trailing_blank_rows(tmp_path):
    sheets = _standard_sheets()
    sheets["Gasoline"].append((None, None, None, None, None))
    path = _make_workbook(tmp_path / "mso.xlsx", sheets)

    result = parse_latest_reserves(path)

    assert result["as_of"] == "2026-09-01"
    assert result["fuels"][0]["days"] == 43


def test_parse_latest_reserves_rejects_mismatched_dates(tmp_path):
    sheets = _standard_sheets()
    sheets["Kerosene"].pop()  # kerosene sheet lags a week
    path = _make_workbook(tmp_path / "mso.xlsx", sheets)

    with pytest.raises(ValueError, match="date"):
        parse_latest_reserves(path)


def test_parse_latest_reserves_rejects_non_numeric_days(tmp_path):
    sheets = _standard_sheets()
    sheets["Diesel"][-1] = (datetime(2026, 9, 1), 2742, 2225, 3039, "-")
    path = _make_workbook(tmp_path / "mso.xlsx", sheets)

    with pytest.raises(ValueError, match="Diesel"):
        parse_latest_reserves(path)


def test_parse_latest_reserves_rejects_implausible_days(tmp_path):
    sheets = _standard_sheets()
    sheets["Gasoline"][-1] = (datetime(2026, 9, 1), 1067, 862, 1823, 1823)
    path = _make_workbook(tmp_path / "mso.xlsx", sheets)

    with pytest.raises(ValueError, match="Gasoline"):
        parse_latest_reserves(path)


def test_parse_latest_reserves_rejects_missing_sheet(tmp_path):
    sheets = _standard_sheets()
    del sheets["Kerosene"]
    path = _make_workbook(tmp_path / "mso.xlsx", sheets)

    with pytest.raises(KeyError):
        parse_latest_reserves(path)


def test_build_reserves_json_matches_existing_file_shape():
    parsed = {
        "as_of": "2026-09-01",
        "fuels": [{"key": "petrol", "label": "Petrol", "days": 43}],
    }

    result = build_reserves_json(parsed)

    assert result == {
        "source": "DCCEEW Minimum Stockholding Obligation",
        "source_url": (
            "https://www.dcceew.gov.au/energy/security/australias-fuel-security/"
            "minimum-stockholding-obligation/statistics"
        ),
        "as_of": "2026-09-01",
        "fuels": [{"key": "petrol", "label": "Petrol", "days": 43}],
    }
    json.dumps(result)  # must be serialisable as-is
