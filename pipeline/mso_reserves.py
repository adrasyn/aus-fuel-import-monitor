"""Parse the DCCEEW MSO weekly snapshot spreadsheet into mso-reserves.json."""

from datetime import datetime

import requests
from openpyxl import load_workbook

SPREADSHEET_URL = (
    "https://www.dcceew.gov.au/sites/default/files/documents/"
    "mso-weekly-snapshot-timeseries.xlsx"
)
SOURCE_URL = (
    "https://www.dcceew.gov.au/energy/security/australias-fuel-security/"
    "minimum-stockholding-obligation/statistics"
)

# (sheet name in workbook, key/label used by the site). Order is render order.
FUEL_SHEETS = [
    ("Gasoline", "petrol", "Petrol"),
    ("Kerosene", "kerosene", "Kerosene"),
    ("Diesel", "diesel", "Diesel"),
]

# Sanity bounds on "days equivalent" — guards against picking up a volume
# column (hundreds/thousands of ML) if the sheet layout shifts.
MIN_DAYS, MAX_DAYS = 10, 90


def download_spreadsheet(output_path: str) -> str:
    resp = requests.get(SPREADSHEET_URL, timeout=60)
    resp.raise_for_status()
    with open(output_path, "wb") as f:
        f.write(resp.content)
    return output_path


def _latest_row(ws) -> tuple:
    """Last row whose first cell is an obligation date. Days equivalent is
    always the final column, whether or not the sheet has the s16A column."""
    latest = None
    for row in ws.iter_rows(min_row=3, values_only=True):
        if isinstance(row[0], datetime):
            latest = row
    if latest is None:
        raise ValueError(f"{ws.title}: no dated rows found")
    return latest


def parse_latest_reserves(excel_path: str) -> dict:
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    as_of = None
    fuels = []
    for sheet, key, label in FUEL_SHEETS:
        row = _latest_row(wb[sheet])
        date = row[0].date().isoformat()
        if as_of is None:
            as_of = date
        elif date != as_of:
            raise ValueError(f"{sheet}: latest date {date} != {as_of}")
        days = row[-1]
        if not isinstance(days, int) or not MIN_DAYS <= days <= MAX_DAYS:
            raise ValueError(f"{sheet}: implausible days value {days!r}")
        fuels.append({"key": key, "label": label, "days": days})
    return {"as_of": as_of, "fuels": fuels}


def build_reserves_json(parsed: dict) -> dict:
    return {
        "source": "DCCEEW Minimum Stockholding Obligation",
        "source_url": SOURCE_URL,
        "as_of": parsed["as_of"],
        "fuels": parsed["fuels"],
    }
